"""生成HomesteadLM论文Word文档"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re

def create_docx(md_path, docx_path, title_style="标题"):
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        
        # 跳过空行但保持间距
        if not line.strip():
            doc.add_paragraph()
            i += 1
            continue
        
        # 标题
        if line.startswith('# ') and not line.startswith('## '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
        
        # 分隔线
        elif line == '---':
            doc.add_paragraph('_' * 40)
        
        # 表格
        elif line.startswith('|'):
            # 收集表格行
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                if '---' not in lines[i]:  # 跳过分隔行
                    table_lines.append(lines[i].strip())
                i += 1
            
            if table_lines:
                # 创建表格
                num_cols = len(table_lines[0].split('|')) - 2
                table = doc.add_table(rows=len(table_lines), cols=num_cols)
                table.style = 'Table Grid'
                
                for row_idx, row_line in enumerate(table_lines):
                    cells = row_line.split('|')[1:-1]  # 去掉首尾的|
                    for col_idx, cell_text in enumerate(cells):
                        cell = table.rows[row_idx].cells[col_idx]
                        cell.text = cell_text.strip()
                continue
        
        # 代码块
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            # 添加为引用格式
            p = doc.add_paragraph()
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.3)
            continue
        
        # 列表
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(line[line.index('. ')+2:], style='List Number')
        
        # 普通文本 - 处理粗体
        else:
            parts = re.split(r'\*\*([^*]+)\*\*', line)
            if len(parts) > 1:
                p = doc.add_paragraph()
                for idx, part in enumerate(parts):
                    run = p.add_run(part)
                    if idx % 2 == 1:  # 粗体部分
                        run.bold = True
            else:
                doc.add_paragraph(line)
        
        i += 1
    
    doc.save(docx_path)
    print(f"已生成: {docx_path}")

# 生成中文版
create_docx(
    "d:/AI研发/LLM研发设想/HomesteadLM-157M论文_中文版.md",
    "d:/AI研发/LLM研发设想/HomesteadLM-157M论文_中文版.docx"
)

# 生成英文版
create_docx(
    "d:/AI研发/LLM研发设想/HomesteadLM-157M论文.md",
    "d:/AI研发/LLM研发设想/HomesteadLM-157M论文.docx"
)

print("完成!")
